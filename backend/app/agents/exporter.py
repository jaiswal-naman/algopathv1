"""
ExporterAgent - The Scribe
Converts LFA JSON data into downloadable documents (DOCX, CSV).
"""

import csv
import json
from io import StringIO, BytesIO
from typing import Dict, Any, List
# We'll need python-docx for Word documents
# pip install python-docx

class Exporter:
    """
    Handles export of LFA data to various formats.
    """

    def __init__(self):
        pass

    def generate_csv(self, lfa_data: Dict[str, Any]) -> str:
        """
        Convert LFA to CSV string.
        Structure: Level, Description, Indicators, Verification, Assumptions

        LFA structure is hierarchical:
        - Goal at root with goal_indicators and assumptions
        - Outcomes contain Outputs
        - Outputs contain Activities
        """
        output = StringIO()
        writer = csv.writer(output)

        # Header
        writer.writerow(["Level", "ID", "Description", "Indicators", "Means of Verification", "Assumptions"])

        # Goal (Impact)
        goal = lfa_data.get("goal", "")
        goal_indicators = "; ".join(lfa_data.get("goal_indicators", []))
        assumptions = "; ".join(lfa_data.get("assumptions", []))
        writer.writerow(["Goal", "", goal, goal_indicators, "", assumptions])

        # Iterate through hierarchical structure: Outcomes -> Outputs -> Activities
        for outcome in lfa_data.get("outcomes", []):
            outcome_id = outcome.get("id", "")
            outcome_desc = outcome.get("description", "")
            outcome_indicators = "; ".join(outcome.get("indicators", []))
            outcome_verification = "; ".join(outcome.get("means_of_verification", []))
            writer.writerow(["Outcome", outcome_id, outcome_desc, outcome_indicators, outcome_verification, ""])

            # Outputs within this Outcome
            for output_item in outcome.get("outputs", []):
                output_id = output_item.get("id", "")
                output_desc = output_item.get("description", "")
                output_indicators = "; ".join(output_item.get("indicators", []))
                output_verification = "; ".join(output_item.get("means_of_verification", []))
                writer.writerow(["Output", output_id, output_desc, output_indicators, output_verification, ""])

                # Activities within this Output
                for activity in output_item.get("activities", []):
                    activity_id = activity.get("id", "")
                    activity_desc = activity.get("description", "")
                    activity_indicators = "; ".join(activity.get("indicators", []))
                    activity_verification = "; ".join(activity.get("means_of_verification", []))
                    writer.writerow(["Activity", activity_id, activity_desc, activity_indicators, activity_verification, ""])

        return output.getvalue()

    def generate_docx(self, lfa_data: Dict[str, Any], metadata: Dict[str, Any]) -> BytesIO:
        """
        Convert LFA to Word Document (BytesIO).
        """
        # Note: Requires 'python-docx' library. 
        # Since we might not have it installed yet, we'll return a placeholder or need to install it.
        # For now, I will assume we can install it.
        
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.table import WD_TABLE_ALIGNMENT
        except ImportError:
            raise ImportError("python-docx is needed for Word export.")

        document = Document()
        
        # Title
        title = metadata.get("title", "Logical Framework Matrix")
        document.add_heading(title, 0)

        # Metadata
        p = document.add_paragraph()
        p.add_run(f"Organization: {metadata.get('organization', 'N/A')}\n").bold = True
        p.add_run(f"Donor: {metadata.get('donor', 'N/A')}\n").bold = True
        p.add_run(f"Date: {metadata.get('date', 'N/A')}")
        
        document.add_heading('Logical Framework Matrix', level=1)

        # Table
        table = document.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        
        # Header Row
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Project Description'
        hdr_cells[1].text = 'Indicators'
        hdr_cells[2].text = 'Means of Verification'
        hdr_cells[3].text = 'Assumptions'

        # Goal
        row = table.add_row().cells
        row[0].text = f"GOAL: {lfa_data.get('goal', '')}"
        row[1].text = "\n".join(lfa_data.get("goal_indicators", []))
        row[3].text = "\n".join(lfa_data.get("assumptions", []))

        # Iterate through hierarchical structure: Outcomes -> Outputs -> Activities
        for outcome in lfa_data.get("outcomes", []):
            row = table.add_row().cells
            row[0].text = f"OUTCOME ({outcome.get('id', '')}): {outcome.get('description', '')}"
            row[1].text = "\n".join(outcome.get("indicators", []))
            row[2].text = "\n".join(outcome.get("means_of_verification", []))

            # Outputs within this Outcome
            for output_item in outcome.get("outputs", []):
                row = table.add_row().cells
                row[0].text = f"  OUTPUT ({output_item.get('id', '')}): {output_item.get('description', '')}"
                row[1].text = "\n".join(output_item.get("indicators", []))
                row[2].text = "\n".join(output_item.get("means_of_verification", []))

                # Activities within this Output
                for activity in output_item.get("activities", []):
                    row = table.add_row().cells
                    row[0].text = f"    ACTIVITY ({activity.get('id', '')}): {activity.get('description', '')}"
                    row[1].text = "\n".join(activity.get("indicators", []))
                    row[2].text = "\n".join(activity.get("means_of_verification", []))
        
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer
